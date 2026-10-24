import pytest
import docx
from backend.resource_processing.pdf_handler import extract_pdf
from backend.text_processing.metadata import build_chunks_from_pdf


def test_build_chunks_from_pdf_basic_flow():
    pdf_result = extract_pdf("tests/TableOfContents.pdf", resource_id="R001")
    chunks = build_chunks_from_pdf(pdf_result)

    assert len(chunks) > 0

    first_chunk = chunks[0]
    assert "chunk_id" in first_chunk
    assert "text" in first_chunk
    assert "metadata" in first_chunk


def test_build_chunks_from_pdf_metadata_schema():
    pdf_result = extract_pdf("tests/TableOfContents.pdf", resource_id="R001")
    chunks = build_chunks_from_pdf(pdf_result)

    required_metadata_keys = {
        "resource_id", "source_type", "source_name", "page_number",
        "chunk_index", "content_type", "word_count", "is_ocr"
    }

    for chunk in chunks:
        assert required_metadata_keys.issubset(chunk["metadata"].keys())


def test_build_chunks_from_pdf_resource_id_propagates():
    pdf_result = extract_pdf("tests/TableOfContents.pdf", resource_id="R001")
    chunks = build_chunks_from_pdf(pdf_result)

    for chunk in chunks:
        assert chunk["metadata"]["resource_id"] == "R001"
        assert chunk["chunk_id"].startswith("R001_C")


def test_build_chunks_from_pdf_chunk_ids_are_unique():
    pdf_result = extract_pdf("tests/TableOfContents.pdf", resource_id="R001")
    chunks = build_chunks_from_pdf(pdf_result)

    chunk_ids = [chunk["chunk_id"] for chunk in chunks]
    assert len(chunk_ids) == len(set(chunk_ids))


def test_build_chunks_from_pdf_chunk_index_is_sequential_across_pages():
    pdf_result = extract_pdf("tests/TableOfContents.pdf", resource_id="R001")
    chunks = build_chunks_from_pdf(pdf_result)

    for expected_index, chunk in enumerate(chunks):
        assert chunk["metadata"]["chunk_index"] == expected_index


def test_build_chunks_from_pdf_is_ocr_defaults_false():
    pdf_result = extract_pdf("tests/TableOfContents.pdf", resource_id="R001")
    chunks = build_chunks_from_pdf(pdf_result)

    for chunk in chunks:
        assert chunk["metadata"]["is_ocr"] is False


def test_build_chunks_from_pdf_word_count_matches_actual_text():
    pdf_result = extract_pdf("tests/TableOfContents.pdf", resource_id="R001")
    chunks = build_chunks_from_pdf(pdf_result)

    for chunk in chunks:
        assert chunk["metadata"]["word_count"] == len(chunk["text"].split())


def test_build_chunks_from_pdf_missing_required_key_raises_error():
    incomplete_result = {
        "resource_id": "R001",
        "source_type": "pdf",
        # "source_name" missing on purpose
        "pages": []
    }

    with pytest.raises(ValueError):
        build_chunks_from_pdf(incomplete_result)


def test_build_chunks_from_pdf_missing_resource_id_raises_error():
    # extract_pdf() allows resource_id=None, but build_chunks_from_pdf()
    # should reject it since chunks need a resource_id downstream.
    pdf_result = extract_pdf("tests/TableOfContents.pdf")  # no resource_id passed

    with pytest.raises(ValueError):
        build_chunks_from_pdf(pdf_result)


def test_build_chunks_from_pdf_empty_resource_id_raises_error():
    pdf_result = extract_pdf("tests/TableOfContents.pdf", resource_id="")

    with pytest.raises(ValueError):
        build_chunks_from_pdf(pdf_result)


from backend.resource_processing.docx_handler import extract_docx
from backend.text_processing.metadata import build_chunks_from_docx


def test_build_chunks_from_docx_basic_flow(tmp_path):
    docx_path = tmp_path / "sample.docx"
    document = docx.Document()
    document.add_heading("Test Document", level=1)
    document.add_paragraph("word " * 300)  # forces multiple chunks
    document.save(str(docx_path))

    docx_result = extract_docx(str(docx_path), resource_id="R002")
    chunks = build_chunks_from_docx(docx_result)

    assert len(chunks) > 0
    for chunk in chunks:
        assert chunk["metadata"]["page_number"] is None
        assert chunk["metadata"]["source_type"] == "docx"
        assert chunk["metadata"]["resource_id"] == "R002"
        assert chunk["chunk_id"].startswith("R002_C")


def test_build_chunks_from_docx_missing_resource_id_raises_error(tmp_path):
    docx_path = tmp_path / "sample.docx"
    document = docx.Document()
    document.add_paragraph("Some content here.")
    document.save(str(docx_path))

    docx_result = extract_docx(str(docx_path))  # no resource_id

    with pytest.raises(ValueError):
        build_chunks_from_docx(docx_result)


def test_build_chunks_from_docx_table_gets_own_chunk(tmp_path):
    docx_path = tmp_path / "sample.docx"
    document = docx.Document()
    document.add_paragraph("Some intro text before the table.")

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Score"
    table.cell(1, 0).text = "Alice"
    table.cell(1, 1).text = "90"

    document.add_paragraph("Some text after the table.")
    document.save(str(docx_path))

    docx_result = extract_docx(str(docx_path), resource_id="R003")
    chunks = build_chunks_from_docx(docx_result)

    table_chunks = [c for c in chunks if c["metadata"]["content_type"] == "table"]
    text_chunks = [c for c in chunks if c["metadata"]["content_type"] == "text"]

    assert len(table_chunks) == 1
    assert len(text_chunks) >= 1
    assert "Alice" in table_chunks[0]["text"]