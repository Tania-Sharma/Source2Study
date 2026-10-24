import pytest
import docx

from backend.resource_processing.docx_handler import extract_docx


def _create_sample_docx(path, include_table=True):
    """Helper: builds a real .docx file with headings, paragraphs,
    a bullet list, and optionally a table -- in a specific order,
    so we can verify iter_inner_content() ordering."""
    document = docx.Document()
    document.add_heading("Chapter 1: Introduction", level=1)
    document.add_paragraph("This is the first paragraph of content.")
    document.add_heading("Section 1.1", level=2)
    document.add_paragraph("Some more explanatory text goes here.")
    document.add_paragraph("First bullet point", style="List Bullet")
    document.add_paragraph("Second bullet point", style="List Bullet")

    if include_table:
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Name"
        table.cell(0, 1).text = "Score"
        table.cell(1, 0).text = "Alice"
        table.cell(1, 1).text = "90"
        document.add_paragraph("Text that comes after the table.")

    document.save(str(path))


def test_extract_docx_basic_flow(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _create_sample_docx(docx_path)

    result = extract_docx(str(docx_path))

    assert result["source_type"] == "docx"
    assert result["element_count"] > 0
    assert len(result["elements"]) == result["element_count"]


def test_extract_docx_captures_headings_and_styles(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _create_sample_docx(docx_path)

    result = extract_docx(str(docx_path))

    paragraph_styles = [
        e["style"] for e in result["elements"] if e["element_type"] == "paragraph"
    ]
    assert "Heading 1" in paragraph_styles
    assert "Heading 2" in paragraph_styles
    assert "List Bullet" in paragraph_styles


def test_extract_docx_element_indices_are_sequential(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _create_sample_docx(docx_path)

    result = extract_docx(str(docx_path))

    for expected_index, element in enumerate(result["elements"], start=1):
        assert element["element_index"] == expected_index


def test_extract_docx_resource_id_propagates(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _create_sample_docx(docx_path)

    result = extract_docx(str(docx_path), resource_id="R002")

    assert result["resource_id"] == "R002"


def test_extract_docx_wrong_extension(tmp_path):
    fake_file = tmp_path / "notes.txt"
    fake_file.write_text("not a real docx")

    with pytest.raises(ValueError):
        extract_docx(str(fake_file))


def test_extract_docx_file_not_found():
    with pytest.raises(FileNotFoundError):
        extract_docx("tests/does_not_exist.docx")


def test_extract_docx_corrupted_file_raises_value_error(tmp_path):
    fake_docx = tmp_path / "corrupted.docx"
    fake_docx.write_bytes(b"this is not a real docx file at all")

    with pytest.raises(ValueError):
        extract_docx(str(fake_docx))


def test_extract_docx_empty_document_raises_value_error(tmp_path):
    empty_docx = tmp_path / "empty.docx"
    document = docx.Document()  # blank document, zero content
    document.save(str(empty_docx))

    with pytest.raises(ValueError):
        extract_docx(str(empty_docx))


def test_extract_docx_table_is_captured(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _create_sample_docx(docx_path, include_table=True)

    result = extract_docx(str(docx_path))

    table_elements = [e for e in result["elements"] if e["element_type"] == "table"]
    assert len(table_elements) == 1

    table = table_elements[0]
    assert table["row_count"] == 2
    assert table["column_count"] == 2
    assert "Name" in table["text"]
    assert "Alice" in table["text"]
    assert "90" in table["text"]


def test_extract_docx_table_preserves_order_with_paragraphs(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _create_sample_docx(docx_path, include_table=True)

    result = extract_docx(str(docx_path))

    types_in_order = [e["element_type"] for e in result["elements"]]

    # The table should appear somewhere after the bullet points and
    # before the final "Text that comes after the table." paragraph.
    table_position = types_in_order.index("table")
    assert types_in_order[table_position - 1] == "paragraph"
    assert types_in_order[table_position + 1] == "paragraph"