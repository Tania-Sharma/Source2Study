import io
import csv
from typing import List, Dict, Any
import openpyxl
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors


def export_to_csv(questions: List[Dict[str, Any]]) -> str:
    """
    Export question bank to CSV format.
    """
    output = io.StringIO()
    writer = csv.writer(output)

    # Header
    writer.writerow([
        "Question ID",
        "Question Text",
        "Subject / Topic",
        "Marks",
        "Section",
        "Occurrences Count",
        "Source Papers & Numbers",
    ])

    for q in questions:
        occ_strings = [
            f"{occ.get('paper_name')} ({occ.get('original_number', 'N/A')})"
            for occ in q.get("occurrences", [])
        ]
        writer.writerow([
            q.get("id"),
            q.get("canonical_text"),
            q.get("topic") or q.get("subject") or "General",
            q.get("marks") or "N/A",
            q.get("section") or "N/A",
            len(q.get("occurrences", [])),
            "; ".join(occ_strings),
        ])

    return output.getvalue()


def export_to_excel(questions: List[Dict[str, Any]]) -> bytes:
    """
    Export question bank to Excel (.xlsx) format using openpyxl.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Question Bank"

    headers = [
        "Question ID",
        "Question Text",
        "Subject / Topic",
        "Marks",
        "Section",
        "Occurrences Count",
        "Source Papers",
    ]
    ws.append(headers)

    # Style header row
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        cell.fill = openpyxl.styles.PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")

    for q in questions:
        occ_strings = [
            f"{occ.get('paper_name')} [{occ.get('original_number', '')}]"
            for occ in q.get("occurrences", [])
        ]
        ws.append([
            q.get("id"),
            q.get("canonical_text"),
            q.get("topic") or q.get("subject") or "General",
            q.get("marks") or "N/A",
            q.get("section") or "N/A",
            len(q.get("occurrences", [])),
            ", ".join(occ_strings),
        ])

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 60)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def export_to_pdf(questions: List[Dict[str, Any]], title: str = "Source2Study Unified Question Bank") -> bytes:
    """
    Export question bank to formatted PDF using reportlab.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()

    elements = []

    # Title
    title_style = ParagraphStyle(
        "TitleStyle",
        parent=styles["Heading1"],
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#1E3A8A"),
        spaceAfter=12,
    )
    elements.append(Paragraph(title, title_style))
    elements.append(Paragraph(f"Consolidated Question Bank &mdash; {len(questions)} unique questions", styles["Normal"]))
    elements.append(Spacer(1, 16))

    body_style = ParagraphStyle(
        "BodyStyle",
        parent=styles["Normal"],
        fontSize=10,
        leading=13,
    )
    meta_style = ParagraphStyle(
        "MetaStyle",
        parent=styles["Normal"],
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#4B5563"),
    )

    table_data = [["#", "Question", "Marks", "Source Papers"]]

    for idx, q in enumerate(questions, start=1):
        occ_strings = [
            f"{occ.get('paper_name')} ({occ.get('original_number', '')})"
            for occ in q.get("occurrences", [])
        ]
        q_para = Paragraph(f"<b>{q.get('canonical_text')}</b>", body_style)
        source_para = Paragraph("<br/>".join(occ_strings), meta_style)

        table_data.append([
            str(idx),
            q_para,
            q.get("marks") or "-",
            source_para,
        ])

    col_widths = [30, 320, 50, 140]
    t = Table(table_data, colWidths=col_widths, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 10),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
        ("TOPPADDING", (0, 0), (-1, 0), 8),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9FAFB")]),
    ]))

    elements.append(t)
    doc.build(elements)
    return buffer.getvalue()
